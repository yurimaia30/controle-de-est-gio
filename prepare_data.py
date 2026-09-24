import json, pathlib, datetime, re, shutil
import openpyxl, pdfplumber

root = pathlib.Path(__file__).parent
(root / 'private').mkdir(exist_ok=True)
(root / 'public').mkdir(exist_ok=True)
if not (root / 'private/index-original.html').exists():
    shutil.copyfile(root / 'index.html', root / 'private/index-original.html')
shutil.copyfile(root / 'FICHA DE FREQUENCIA DE RADIOLOGIA.pdf', root / 'public/modelo-frequencia.pdf')
shutil.copyfile(root / 'logo grau.jpg', root / 'public/logo.jpg')
def txt(v): return str(v).strip() if v is not None else ''
def date(v): return v.date().isoformat() if isinstance(v, datetime.datetime) else ''
week = ['Segunda','Terça','Quarta','Quinta','Sexta']
w = openpyxl.load_workbook(root / 'Planilha de Estágio - Radiologia.xlsx', data_only=True)
fields = []
for key, turno, prec in [('IOT','Tarde','Equipe Técnica IOT'),('SESI','Manhã','Equipe Técnica SESI')]:
    slots=[]
    for r in list(w[key].values)[1:]:
        if r[0] != key or not isinstance(r[2], (int,float)): continue
        for i,d in enumerate(week):
            value=txt(r[i+3])
            slots.append(dict(id=f'{int(r[2])}-{i+1}',vaga=int(r[2]),dias=[i+1],aluno=value if value!='X' else '',bloqueada=value=='X',inicio='',fim='',obs=''))
    fields.append(dict(id=key,group=key,nome=key,turno=turno,preceptor=prec,slots=slots,observacoes='',horas=None))
upa=w['UPA']
for n,rows,col,days,turn in [(1,range(6,11),2,[1,3,5],'Noite'),(2,range(6,11),7,[2,4],'Noite'),(3,range(14,19),2,[6],'Tarde e noite'),(4,range(14,19),7,[0],'Manhã, tarde e noite')]:
    slots=[dict(id=str(int(upa.cell(r,col).value)),vaga=int(upa.cell(r,col).value),dias=days,aluno=txt(upa.cell(r,col+1).value),inicio=date(upa.cell(r,col+2).value),fim=date(upa.cell(r,col+3).value),bloqueada=False,obs='') for r in rows]
    fields.append(dict(id=f'UPA-G{n}',group='UPA',nome=f'UPA Grupo {n}',turno=turn,preceptor='Erick Carvalho' if n==4 else 'Francisco Daniel',slots=slots,observacoes='',horas=None))
fields.append(dict(id='UPA-MANHA',group='UPA',nome='UPA Manhã',turno='Manhã',preceptor='Artur Nogueira',slots=[],observacoes='Quantidade de vagas não informada na planilha. Cadastrar após confirmação.',horas=None))
students=[]
hours=openpyxl.load_workbook(root / 'CONTROLE DE HORAS.xlsx',data_only=True).active
for rownum,r in enumerate(hours.values,1):
    if rownum<3 or not isinstance(r[0],str) or not r[0].strip(): continue
    hist=[]
    for local,period,h in [(2,3,6),(7,8,9),(10,11,12),(13,14,15),(16,17,18)]:
        if r[local] is not None:
            hist.append(dict(local=txt(r[local]),periodo=txt(r[period]),horas=r[h] if isinstance(r[h],(int,float)) else 0))
    students.append(dict(id=f'planilha-linha-{rownum}',nome=txt(r[0]),turma='',historico=hist,horasPlanilha=r[19] if isinstance(r[19],(int,float)) else None,fonte=f'CONTROLE DE HORAS.xlsx / Planilha1 / linha {rownum}'))
wait=[]
for r in list(w['Lista de espera '].values)[1:]:
    if r[0]: wait.append(dict(nome=txt(r[0]),turma=txt(r[1]),disponibilidade=txt(r[2]),campo=txt(r[3])))
for r in range(6,12):
    wait.append(dict(nome=txt(upa.cell(r,12).value),turma=txt(upa.cell(r,13).value),disponibilidade='Noite e final de semana',campo=txt(upa.cell(r,14).value)))
payload=dict(version=2,fields=fields,students=students,waitlist=wait)
(root/'private/importacao-planilhas.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf8')
p=pdfplumber.open(root/'FICHA DE FREQUENCIA DE RADIOLOGIA.pdf').pages[0]
ys=sorted(set(round(r['top'],2) for r in p.rects if abs(r['x0']-30.96)<.1 and r['height']>10))
(root/'public/pdf-layout.json').write_text(json.dumps(dict(rows=[y for y in ys if 202<=y<620])),encoding='utf8')
print(f'{len(students)} alunos; {len(fields)} campos; {len(wait)} espera; linhas PDF: {ys}')
